"""Applique les corrections documentaires auditees a la banque de 100 questions.

Le document Word est corrige sans reconstruire sa mise en page. Le script est
idempotent afin de pouvoir etre rejoue sans effet cumulatif.
"""

from __future__ import annotations

import os
import re
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT_PATH = ROOT / "data" / "questions" / "Banque_100_questions_CORRIGEE.docx"

QUESTION_UPDATES = {
    9: "Combien de minutes fallait-il en moyenne, en 2021, pour atteindre un établissement sanitaire de premier contact ?",
    10: "Quel était, en 2021, le temps moyen nécessaire pour atteindre un établissement sanitaire de niveau secondaire ?",
    11: "Quel était, en 2021, le temps moyen nécessaire pour atteindre un CHU ou un institut spécialisé ?",
    17: "Quel était le taux de mortalité maternelle en 2021 ?",
    31: "Combien d’entreprises formelles étaient recensées en 2024 ?",
    32: "Quelle était la proportion de PME parmi les entreprises formelles en 2024 ?",
    33: "Quel était l’Excédent Brut d’Exploitation des entreprises formelles en 2022 ?",
    41: "Quel était le taux de chômage global en 2021 ?",
    42: "Quel était le taux de chômage des femmes en 2021 ?",
    43: "Quel était le taux de chômage à Abidjan en 2021 ?",
    44: "Quel était le taux de pauvreté national en 2021 ?",
    45: "Quel était le taux de pauvreté en milieu rural en 2021 ?",
    46: "Quel était le taux de pauvreté en milieu urbain en 2021 ?",
    47: "Quel était l’indice de Gini national en 2021 ?",
    48: "Quelle part de la consommation annuelle revenait aux 10 % les plus pauvres en 2021 ?",
    49: "Quelle part de la consommation annuelle revenait aux 10 % les plus riches en 2021 ?",
    72: "Quel était le PIB en volume CVS de la Côte d’Ivoire au quatrième trimestre 2025, aux prix chaînés de 2015 ?",
    73: "De combien le PIB en volume CVS a-t-il varié entre le troisième et le quatrième trimestre 2025 ?",
    74: "Quelle était la valeur ajoutée nominale totale de l’agriculture vivrière sur l’ensemble de l’année 2025 ?",
    77: "En rapportant le déficit budgétaire 2025 au PIB nominal annuel 2025, quelle part du PIB représentait-il ?",
    99: "Quel a été le taux de croissance des importations ivoiriennes sur les douze mois de 2025 par rapport à 2024 ?",
}

ANSWER_UPDATES = {
    12: "94,3 % (estimation).",
    13: "98,6 % (estimation).",
    14: "70,8 % (estimation).",
    22: "1 556 000 FCFA par habitant (soit 1 556 milliers de FCFA).",
    29: "-0,4 %.",
    30: "6,3 %.",
    72: "12 083,3 milliards de FCFA.",
    73: "Une baisse d’environ 0,37 %.",
    74: "3 357,6 milliards de FCFA.",
    77: "Environ 2,90 % du PIB.",
    99: "7,93 %.",
}

SOURCE_CNT_14 = (
    "data/sources/eco_rci → ANSTAT → Comptes_Nationaux → "
    "Comptes_nationaux_trimestriels (1er trimestriels 2026).pdf → page 14"
)
SOURCE_CNT_17 = SOURCE_CNT_14.rsplit("page 14", 1)[0] + "page 17"
SOURCE_TBSE_12 = (
    "data/sources/eco_rci → ANSTAT → "
    "TABLEAU DE BORD SOCIO-ECONOMIQUE (TBSE) MARS 2026 → "
    "TABLEAU DE BORD SOCIO-ECONOMIQUE (TBSE) DECEMBRE 2025.pdf → page 12"
)
SOURCE_TBSE_13 = SOURCE_TBSE_12.rsplit("page 12", 1)[0] + "page 13"

SOURCE_UPDATES = {
    72: SOURCE_CNT_14,
    73: SOURCE_CNT_14,
    74: SOURCE_CNT_17,
    77: f"{SOURCE_TBSE_12} || {SOURCE_CNT_17}",
    99: SOURCE_TBSE_13,
}

ANSWER_MODE_UPDATES = {
    28: "Calculée",
    72: "Factuelle",
    73: "Calculée",
    74: "Calculée",
    77: "Calculée multi-source",
    99: "Factuelle",
}

QUESTION_RE = re.compile(r"^(\d+)\.\s+(.+?)\s*$")


def clean(text: str) -> str:
    return " ".join(str(text).replace("\xa0", " ").split())


def current_locator(locator: str) -> str:
    value = locator.replace("ECO_RCI_stage_final", "data/sources/eco_rci")
    value = re.sub(
        r"(?<![/\\])([A-Za-z0-9_]+_page)\.txt(?=\s*→\s*Publication)",
        r"data/facebook/raw/\1.jsonl",
        value,
        flags=re.IGNORECASE,
    )
    return value


def find_metadata_table(document: Document):
    expected = ["n°", "domaine", "statut", "source / chemin", "réponse"]
    for table in document.tables:
        if not table.rows:
            continue
        header = [clean(cell.text).casefold() for cell in table.rows[0].cells]
        if header == expected:
            return table
    raise ValueError("Tableau récapitulatif des 100 questions introuvable")


def set_runs(paragraph, segments) -> None:
    """Remplace le texte en reutilisant la mise en forme de runs modeles."""
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    for text, template_run in segments:
        run = paragraph.add_run(text)
        if template_run._r.rPr is not None:
            run._r.insert(0, deepcopy(template_run._r.rPr))


def update_document(document: Document) -> None:
    question_paragraphs: dict[int, int] = {}
    for index, paragraph in enumerate(document.paragraphs):
        match = QUESTION_RE.match(clean(paragraph.text))
        if match:
            question_paragraphs[int(match.group(1))] = index
    if set(question_paragraphs) != set(range(1, 101)):
        raise ValueError("Le document ne contient pas exactement les questions 1 à 100")

    template_index = question_paragraphs[1]
    question_template_runs = list(document.paragraphs[template_index].runs)
    type_template_run = document.paragraphs[template_index + 1].runs[0]
    answer_template_runs = list(document.paragraphs[template_index + 2].runs)
    source_template_run = document.paragraphs[template_index + 3].runs[0]

    for paragraph in document.paragraphs:
        refreshed = current_locator(paragraph.text)
        if refreshed != paragraph.text:
            paragraph.text = refreshed

    for number, index in question_paragraphs.items():
        if number in QUESTION_UPDATES:
            set_runs(document.paragraphs[index], [
                (f"{number}. ", question_template_runs[0]),
                (QUESTION_UPDATES[number], question_template_runs[1]),
            ])

        type_paragraph = document.paragraphs[index + 1]
        answer_paragraph = document.paragraphs[index + 2]
        source_paragraph = document.paragraphs[index + 3]
        status = "IN-BASE" if number <= 70 or number in SOURCE_UPDATES else "OUT-OF-BASE"

        if number in ANSWER_MODE_UPDATES:
            set_runs(type_paragraph, [(
                f"Type : {ANSWER_MODE_UPDATES[number]}    |    Statut : {status}",
                type_template_run,
            )])
        elif number in SOURCE_UPDATES:
            prefix = type_paragraph.text.split("|", 1)[0].strip()
            set_runs(type_paragraph, [(
                f"{prefix}    |    Statut : {status}", type_template_run,
            )])

        if number in ANSWER_UPDATES:
            set_runs(answer_paragraph, [
                ("Réponse : ", answer_template_runs[0]),
                (ANSWER_UPDATES[number], answer_template_runs[1]),
            ])
        if number in SOURCE_UPDATES:
            set_runs(source_paragraph, [(
                f"Chemin d’accès : {SOURCE_UPDATES[number]}", source_template_run,
            )])

    for paragraph in document.paragraphs[:10]:
        if "questions avec réponse" in paragraph.text:
            paragraph.text = "75 questions avec réponse dans la base • 25 questions hors base"

    # Petites corrections éditoriales sans perdre la mise en forme existante.
    if len(document.paragraphs[1].runs) >= 2:
        document.paragraphs[1].runs[1].text = "réponses"
    for run in document.paragraphs[6].runs:
        if "Pour les publications Facebook" in run.text:
            run.text = (
                "Pour les publications Facebook, le chemin indique le fichier JSONL brut et "
                "le numéro de publication."
            )

    metadata = find_metadata_table(document)
    rows = {int(clean(row.cells[0].text)): row for row in metadata.rows[1:]}
    if set(rows) != set(range(1, 101)):
        raise ValueError("Le tableau de métadonnées n’est pas complet")
    status_template_run = rows[1].cells[2].paragraphs[0].runs[0]
    answer_template_run = rows[1].cells[4].paragraphs[0].runs[0]
    for number, row in rows.items():
        row.cells[3].text = current_locator(row.cells[3].text)
        if number in SOURCE_UPDATES:
            set_runs(
                row.cells[2].paragraphs[0],
                [("IN-BASE", status_template_run)],
            )
            row.cells[3].text = SOURCE_UPDATES[number]
        if number in ANSWER_UPDATES:
            set_runs(
                row.cells[4].paragraphs[0],
                [(ANSWER_UPDATES[number], answer_template_run)],
            )

    document.tables[0].cell(1, 1).text = "75"
    document.tables[0].cell(2, 1).text = "25"
    document.core_properties.modified = datetime.now(timezone.utc)


def validate_document(path: Path) -> None:
    document = Document(path)
    metadata = find_metadata_table(document)
    rows = {int(clean(row.cells[0].text)): row for row in metadata.rows[1:]}
    in_base = sum(clean(row.cells[2].text).upper() == "IN-BASE" for row in rows.values())
    out_of_base = sum(clean(row.cells[2].text).upper() == "OUT-OF-BASE" for row in rows.values())
    if (in_base, out_of_base) != (75, 25):
        raise ValueError(f"Répartition invalide : {in_base} IN-BASE, {out_of_base} OUT-OF-BASE")
    if any("ECO_RCI_stage_final" in paragraph.text for paragraph in document.paragraphs):
        raise ValueError("Un ancien chemin subsiste dans le corps du document")
    if any("ECO_RCI_stage_final" in row.cells[3].text for row in metadata.rows[1:]):
        raise ValueError("Un ancien chemin subsiste dans le tableau récapitulatif")


def main() -> None:
    document = Document(DOCUMENT_PATH)
    update_document(document)
    temporary_path = DOCUMENT_PATH.with_suffix(".corrected.tmp.docx")
    document.save(temporary_path)
    validate_document(temporary_path)
    os.replace(temporary_path, DOCUMENT_PATH)
    print(f"Document corrigé : {DOCUMENT_PATH}")


if __name__ == "__main__":
    main()
