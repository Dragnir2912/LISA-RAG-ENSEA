"""Import reproductible des PDF et classeurs locaux dans le corpus brut du RAG.

Les originaux restent sous ``data/sources``. Ce module produit uniquement des
enregistrements JSON régénérables dans ``data/raw`` ; le nettoyage, le chunking
et l'indexation restent assurés par les composants habituels du pipeline.
"""

from __future__ import annotations

import hashlib
import json
import logging
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Callable
from urllib.parse import quote

from config import Config, load_config
from scraper import url_hash
from table_facts import (
    TABLE_EXTRACTION_VERSION,
    SPREADSHEET_EXTRACTION_VERSION,
    deduplicate_table_facts,
    extract_layout_table_facts,
    extract_matrix_table_facts,
    infer_table_context,
    is_temporal_header_row,
    rotated_words_to_layout,
)

logger = logging.getLogger(__name__)


@dataclass
class LocalPDFImportReport:
    discovered_pdfs: int = 0
    imported_pdfs: int = 0
    reused_pdfs: int = 0
    duplicate_pdfs: int = 0
    failed_pdfs: int = 0
    discovered_spreadsheets: int = 0
    imported_spreadsheets: int = 0
    reused_spreadsheets: int = 0
    duplicate_spreadsheets: int = 0
    failed_spreadsheets: int = 0
    spreadsheet_sheets_detected: int = 0
    ignored_non_pdf_files: int = 0
    removed_stale_records: int = 0
    extracted_characters: int = 0
    table_pages_detected: int = 0
    table_facts_extracted: int = 0
    table_extraction_failures: int = 0
    failures: list[dict[str, str]] = field(default_factory=list)


class LocalPDFImporter:
    """Transforme les PDF et XLSX d'une bibliothèque locale en records compatibles."""

    def __init__(self, config: Config) -> None:
        self.config = config
        self.cfg = config.local_documents
        self.source_dir = self.cfg.source_full_path
        self.raw_dir = self.cfg.raw_full_path
        self.report_path = self.cfg.report_full_path

    @staticmethod
    def _title(path: Path) -> str:
        title = re.sub(r"[_-]+", " ", path.stem)
        return re.sub(r"\s+", " ", title).strip()

    @staticmethod
    def _document_year(path: Path) -> str | None:
        years = re.findall(r"\b(?:19|20)\d{2}\b", path.stem)
        return max(years) if years else None

    @staticmethod
    def _extract_pages(
        path: Path,
        progress_callback: Callable[[int, int, str], None] | None = None,
    ) -> tuple[str, int]:
        from pypdf import PdfReader

        sections = []
        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                reader.decrypt("")
            page_count = len(reader.pages)
            if progress_callback:
                progress_callback(0, page_count, "pypdf")
            for page_number, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    sections.append(f"## Page {page_number}\n\n{text}")
                if progress_callback:
                    progress_callback(page_number, page_count, "pypdf")
            if sections:
                return "\n\n".join(sections), page_count
        except Exception as error:
            logger.warning("pypdf a échoué sur %s : %s", path, error)

        # Certains PDF encodent le texte d'une manière mieux comprise par
        # pdfplumber. Ce second passage n'est effectué qu'en cas d'échec total.
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            page_count = len(pdf.pages)
            if progress_callback:
                progress_callback(0, page_count, "pdfplumber")
            for page_number, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    sections.append(f"## Page {page_number}\n\n{text}")
                if progress_callback:
                    progress_callback(page_number, page_count, "pdfplumber")
            return "\n\n".join(sections), page_count

    def _extract_table_facts(
        self,
        path: Path,
        progress_callback: Callable[[int, int, str], None] | None = None,
    ) -> tuple[list[dict], int, int]:
        """Extrait les series periodiques en preservant les colonnes visuelles.

        Une erreur de mise en page ne doit jamais rendre le PDF inutilisable :
        le texte narratif reste importe et l'echec tabulaire est comptabilise.
        """
        if not self.cfg.extract_table_facts:
            return [], 0, 0

        import pdfplumber

        facts: list[dict] = []
        pages_with_facts = 0
        failed_pages = 0
        with pdfplumber.open(str(path)) as pdf:
            total = len(pdf.pages)
            if progress_callback:
                progress_callback(0, total, "tableaux")
            for page_number, page in enumerate(pdf.pages, start=1):
                try:
                    standard_words = page.extract_words(
                        x_tolerance=1.5,
                        y_tolerance=2.0,
                        use_text_flow=False,
                    )
                    context = infer_table_context(
                        " ".join(str(word.get("text", "")) for word in standard_words),
                        source_hint=str(path),
                    )
                    page_facts = extract_layout_table_facts(
                        standard_words,
                        page_number=page_number,
                        min_periods=self.cfg.table_min_periods,
                        max_facts=self.cfg.table_max_facts_per_page,
                        default_unit=context["unit"],
                        territory=context["territory"],
                        measure_type=context["measure_type"],
                    )

                    # Les tableaux paysages inseres a 90 degres sont frequents
                    # dans les comptes nationaux. On les remet dans leur repere
                    # logique avant d'appliquer exactement le meme extracteur.
                    rotated_chars = [char for char in page.chars if not char.get("upright", True)]
                    if len(rotated_chars) >= max(20, len(page.chars) // 4):
                        counterclockwise = sum(
                            float(char.get("matrix", (0, 0, 0, 0))[1]) > 0
                            and float(char.get("matrix", (0, 0, 0, 0))[2]) < 0
                            for char in rotated_chars
                        ) >= len(rotated_chars) / 2
                        rotated_words = page.extract_words(
                            x_tolerance=1.5,
                            y_tolerance=2.0,
                            use_text_flow=False,
                            line_dir_rotated="ltr",
                            char_dir_rotated=("btt" if counterclockwise else "ttb"),
                        )
                        rotated_context = infer_table_context(
                            " ".join(
                                str(word.get("text", ""))
                                for word in rotated_words if not word.get("upright", True)
                            ),
                            source_hint=str(path),
                        )
                        layout_words = rotated_words_to_layout(
                            rotated_words,
                            page_width=float(page.width),
                            page_height=float(page.height),
                            clockwise=not counterclockwise,
                        )
                        page_facts.extend(extract_layout_table_facts(
                            layout_words,
                            page_number=page_number,
                            min_periods=self.cfg.table_min_periods,
                            max_facts=self.cfg.table_max_facts_per_page,
                            default_unit=rotated_context["unit"],
                            territory=rotated_context["territory"],
                            measure_type=rotated_context["measure_type"],
                        ))
                        page_facts = deduplicate_table_facts(page_facts)
                    if page_facts:
                        pages_with_facts += 1
                        facts.extend(page_facts)
                except Exception as error:
                    failed_pages += 1
                    logger.warning(
                        "Extraction tabulaire ignoree pour %s page %d : %s",
                        path,
                        page_number,
                        error,
                    )
                if progress_callback:
                    progress_callback(page_number, total, "tableaux")
        return deduplicate_table_facts(facts), pages_with_facts, failed_pages

    @staticmethod
    def _excel_cell_text(cell) -> str:
        value = cell.value
        if value is None:
            return ""
        if isinstance(value, (datetime, date)):
            months = (
                "", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin",
                "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre",
            )
            return f"{months[value.month]} {value.year}"
        if isinstance(value, bool):
            return "Oui" if value else "Non"
        if isinstance(value, (int, float)):
            number_format = str(getattr(cell, "number_format", "") or "")
            numeric = float(value)
            if "%" in number_format:
                numeric *= 100.0
                decimals_match = re.search(r"0\.(0+)\s*%", number_format)
                decimals = len(decimals_match.group(1)) if decimals_match else 0
                rendered = f"{numeric:.{decimals}f}"
                return rendered.replace(".", ",") + "%"
            if numeric.is_integer():
                return str(int(numeric))
            return f"{numeric:.6f}".rstrip("0").rstrip(".").replace(".", ",")
        return re.sub(r"\s+", " ", str(value)).strip()

    def _extract_spreadsheet(
        self,
        path: Path,
        progress_callback: Callable[[int, int, str], None] | None = None,
    ) -> tuple[str, list[dict], int, int]:
        """Extrait chaque feuille et segmente ses tableaux à en-têtes répétés."""
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sections: list[str] = []
        facts: list[dict] = []
        failures = 0
        total = len(workbook.worksheets)
        if progress_callback:
            progress_callback(0, total, "tableur")
        try:
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                try:
                    rows: list[list[str]] = []
                    for excel_row in sheet.iter_rows():
                        values = [self._excel_cell_text(cell) for cell in excel_row]
                        while values and not values[-1]:
                            values.pop()
                        rows.append(values)
                    while rows and not any(rows[-1]):
                        rows.pop()
                    if not rows:
                        continue

                    section_lines = [f"## Feuille {sheet_index} : {sheet.title}"]
                    section_lines.extend(
                        " | ".join(cell for cell in row if cell)
                        for row in rows if any(row)
                    )
                    sections.append("\n\n".join(section_lines))

                    header_indices = [
                        index for index, row in enumerate(rows)
                        if is_temporal_header_row(row)
                    ]
                    for table_number, header_index in enumerate(header_indices, start=1):
                        next_header = (
                            header_indices[table_number]
                            if table_number < len(header_indices) else len(rows)
                        )
                        block_start = max(0, header_index - 2)
                        block = rows[block_start:next_header]
                        context_text = " ".join(
                            [path.stem, sheet.title]
                            + [cell for row in block[:4] for cell in row]
                        )
                        context = infer_table_context(
                            context_text,
                            source_hint=str(path),
                        )
                        block_facts = extract_matrix_table_facts(
                            block,
                            page_number=sheet_index,
                            table_number=table_number,
                            default_unit=context["unit"],
                            territory=context["territory"],
                            measure_type=context["measure_type"],
                        )
                        for fact in block_facts:
                            if "%" in str(fact.get("value", "")) and not fact.get("unit"):
                                fact["unit"] = "%"
                            row_context = str(fact.get("qualifier", "")).strip()
                            fact["qualifier"] = " | ".join(filter(None, (
                                f"Feuille : {sheet.title}", row_context,
                            )))
                            fact["source_method"] = "spreadsheet_matrix"
                        facts.extend(block_facts)
                except Exception as error:
                    failures += 1
                    logger.warning(
                        "Feuille tableur ignorée pour %s (%s) : %s",
                        path, sheet.title, error,
                    )
                if progress_callback:
                    progress_callback(sheet_index, total, "tableur")
        finally:
            workbook.close()
        return (
            "\n\n".join(sections),
            deduplicate_table_facts(facts),
            total,
            failures,
        )

    def _record_url(self, relative_path: Path) -> str:
        encoded = quote(relative_path.as_posix(), safe="/")
        return f"{self.cfg.public_route_prefix.rstrip('/')}/{encoded}"

    def _write_report(self, report: LocalPDFImportReport) -> None:
        self.report_path.parent.mkdir(parents=True, exist_ok=True)
        self.report_path.write_text(
            json.dumps(asdict(report), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def import_all(self, *, show_progress: bool = True) -> LocalPDFImportReport:
        if not self.source_dir.exists():
            raise FileNotFoundError(f"Bibliothèque locale absente : {self.source_dir}")

        self.raw_dir.mkdir(parents=True, exist_ok=True)
        all_files = sorted(path for path in self.source_dir.rglob("*") if path.is_file())
        pdf_paths = [path for path in all_files if path.suffix.casefold() == ".pdf"]
        spreadsheet_paths = [
            path for path in all_files if path.suffix.casefold() in {".xlsx", ".xlsm"}
        ]
        report = LocalPDFImportReport(
            discovered_pdfs=len(pdf_paths),
            discovered_spreadsheets=len(spreadsheet_paths),
            ignored_non_pdf_files=(
                len(all_files) - len(pdf_paths) - len(spreadsheet_paths)
            ),
        )
        expected_records: set[Path] = set()
        seen_digests: dict[str, Path] = {}

        if show_progress:
            from tqdm.auto import tqdm

            pdf_iterator = tqdm(
                pdf_paths,
                total=len(pdf_paths),
                desc="Import des PDF ECO RCI",
                unit="PDF",
                dynamic_ncols=True,
            )
        else:
            pdf_iterator = pdf_paths

        for pdf_path in pdf_iterator:
            relative_path = pdf_path.relative_to(self.source_dir)
            if show_progress:
                pdf_iterator.set_description(
                    f"PDF : {pdf_path.name[:42]}",
                    refresh=True,
                )
                pdf_iterator.set_postfix(
                    phase="ouverture",
                    importés=report.imported_pdfs,
                    réutilisés=report.reused_pdfs,
                    erreurs=report.failed_pdfs,
                    refresh=True,
                )
            try:
                digest = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
                if digest in seen_digests:
                    report.duplicate_pdfs += 1
                    logger.warning(
                        "PDF dupliqué ignoré : %s (identique à %s)",
                        relative_path,
                        seen_digests[digest].relative_to(self.source_dir),
                    )
                    continue
                seen_digests[digest] = pdf_path

                url = self._record_url(relative_path)
                output_path = self.raw_dir / f"local_pdf_{url_hash(url)}.json"
                expected_records.add(output_path.resolve())
                if output_path.exists():
                    existing = json.loads(output_path.read_text(encoding="utf-8"))
                    existing_metadata = existing.get("metadata", {})
                    table_version_current = (
                        not self.cfg.extract_table_facts
                        or existing_metadata.get("table_extraction_version")
                        == TABLE_EXTRACTION_VERSION
                    )
                    if (
                        existing_metadata.get("file_sha256") == digest
                        and table_version_current
                    ):
                        report.reused_pdfs += 1
                        report.extracted_characters += len(existing.get("text", ""))
                        existing_facts = existing.get("table_facts", [])
                        report.table_facts_extracted += len(existing_facts)
                        report.table_pages_detected += len({
                            int(fact.get("page") or 0)
                            for fact in existing_facts if fact.get("page")
                        })
                        continue

                def page_progress(current: int, total: int, engine: str) -> None:
                    if show_progress:
                        pdf_iterator.set_postfix(
                            pages=f"{current}/{total}",
                            moteur=engine,
                            importés=report.imported_pdfs,
                            réutilisés=report.reused_pdfs,
                            erreurs=report.failed_pdfs,
                            refresh=True,
                        )

                text, page_count = self._extract_pages(
                    pdf_path,
                    progress_callback=page_progress,
                )
                if len(re.sub(r"\s+", " ", text).strip()) < 80:
                    raise ValueError("aucun texte exploitable ; OCR probablement nécessaire")

                table_facts, table_pages, table_failures = self._extract_table_facts(
                    pdf_path,
                    progress_callback=page_progress,
                )
                report.table_facts_extracted += len(table_facts)
                report.table_pages_detected += table_pages
                report.table_extraction_failures += table_failures

                title = self._title(pdf_path)
                source_parts = list(relative_path.parts[:-1])
                markdown = f"# {title}\n\n{text}"
                record = {
                    "url": url,
                    "title": title,
                    "headings": [{"level": "h1", "text": title}],
                    "text": markdown,
                    "raw_html": "",
                    "lists": [],
                    "tables": [],
                    "table_facts": table_facts,
                    "links": [],
                    "breadcrumbs": source_parts,
                    "metadata": {
                        "doc_type": "pdf",
                        "source": source_parts[0] if source_parts else "eco_rci",
                        "source_path": relative_path.as_posix(),
                        "file_sha256": digest,
                        "page_count": page_count,
                        "table_extraction_version": TABLE_EXTRACTION_VERSION,
                        "table_facts": len(table_facts),
                        "table_pages": table_pages,
                    },
                    "published_date": self._document_year(pdf_path),
                    "modified_date": None,
                    "scraped_at": datetime.now(timezone.utc).isoformat(),
                    "content_hash": hashlib.sha256(
                        (
                            markdown
                            + json.dumps(table_facts, ensure_ascii=False, sort_keys=True)
                        ).encode("utf-8")
                    ).hexdigest(),
                }
                output_path.write_text(
                    json.dumps(record, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                report.imported_pdfs += 1
                report.extracted_characters += len(markdown)
            except Exception as error:  # un PDF illisible ne bloque pas les 79 autres
                report.failed_pdfs += 1
                report.failures.append({
                    "path": relative_path.as_posix(),
                    "error": f"{type(error).__name__}: {error}",
                })
                logger.error(
                    "PDF non importé : %s — %s: %s",
                    relative_path,
                    type(error).__name__,
                    error,
                )
            finally:
                if show_progress:
                    pdf_iterator.set_postfix(
                        importés=report.imported_pdfs,
                        réutilisés=report.reused_pdfs,
                        doublons=report.duplicate_pdfs,
                        erreurs=report.failed_pdfs,
                        refresh=True,
                    )

        if show_progress:
            pdf_iterator.set_description("Import ECO RCI terminé", refresh=True)

        if show_progress:
            from tqdm.auto import tqdm

            spreadsheet_iterator = tqdm(
                spreadsheet_paths,
                total=len(spreadsheet_paths),
                desc="Import des tableurs ECO RCI",
                unit="classeur",
                dynamic_ncols=True,
            )
        else:
            spreadsheet_iterator = spreadsheet_paths

        for spreadsheet_path in spreadsheet_iterator:
            relative_path = spreadsheet_path.relative_to(self.source_dir)
            if show_progress:
                spreadsheet_iterator.set_description(
                    f"Tableur : {spreadsheet_path.name[:38]}", refresh=True,
                )
            try:
                digest = hashlib.sha256(spreadsheet_path.read_bytes()).hexdigest()
                if digest in seen_digests:
                    report.duplicate_spreadsheets += 1
                    logger.warning(
                        "Tableur dupliqué ignoré : %s (identique à %s)",
                        relative_path,
                        seen_digests[digest].relative_to(self.source_dir),
                    )
                    continue
                seen_digests[digest] = spreadsheet_path

                url = self._record_url(relative_path)
                output_path = self.raw_dir / f"local_sheet_{url_hash(url)}.json"
                expected_records.add(output_path.resolve())
                if output_path.exists():
                    existing = json.loads(output_path.read_text(encoding="utf-8"))
                    existing_metadata = existing.get("metadata", {})
                    if (
                        existing_metadata.get("file_sha256") == digest
                        and existing_metadata.get("table_extraction_version")
                        == SPREADSHEET_EXTRACTION_VERSION
                    ):
                        report.reused_spreadsheets += 1
                        report.extracted_characters += len(existing.get("text", ""))
                        existing_facts = existing.get("table_facts", [])
                        report.table_facts_extracted += len(existing_facts)
                        sheet_count = int(existing_metadata.get("sheet_count") or 0)
                        report.spreadsheet_sheets_detected += sheet_count
                        report.table_pages_detected += len({
                            int(fact.get("page") or 0)
                            for fact in existing_facts if fact.get("page")
                        })
                        continue

                def sheet_progress(current: int, total: int, engine: str) -> None:
                    if show_progress:
                        spreadsheet_iterator.set_postfix(
                            feuilles=f"{current}/{total}",
                            moteur=engine,
                            importés=report.imported_spreadsheets,
                            erreurs=report.failed_spreadsheets,
                            refresh=True,
                        )

                text, table_facts, sheet_count, sheet_failures = self._extract_spreadsheet(
                    spreadsheet_path,
                    progress_callback=sheet_progress,
                )
                if len(re.sub(r"\s+", " ", text).strip()) < 40:
                    raise ValueError("aucune feuille exploitable")
                if not table_facts:
                    raise ValueError("aucun tableau périodique exploitable")

                report.table_facts_extracted += len(table_facts)
                report.spreadsheet_sheets_detected += sheet_count
                report.table_pages_detected += len({
                    int(fact.get("page") or 0)
                    for fact in table_facts if fact.get("page")
                })
                report.table_extraction_failures += sheet_failures

                title = self._title(spreadsheet_path)
                source_parts = list(relative_path.parts[:-1])
                markdown = f"# {title}\n\n{text}"
                record = {
                    "url": url,
                    "title": title,
                    "headings": [{"level": "h1", "text": title}],
                    "text": markdown,
                    "raw_html": "",
                    "lists": [],
                    "tables": [],
                    "table_facts": table_facts,
                    "links": [],
                    "breadcrumbs": source_parts,
                    "metadata": {
                        "doc_type": "spreadsheet",
                        "source": source_parts[0] if source_parts else "eco_rci",
                        "source_path": relative_path.as_posix(),
                        "file_sha256": digest,
                        "sheet_count": sheet_count,
                        "table_extraction_version": SPREADSHEET_EXTRACTION_VERSION,
                        "table_facts": len(table_facts),
                        "table_sheets": len({
                            int(fact.get("page") or 0)
                            for fact in table_facts if fact.get("page")
                        }),
                    },
                    "published_date": self._document_year(spreadsheet_path),
                    "modified_date": None,
                    "scraped_at": datetime.now(timezone.utc).isoformat(),
                    "content_hash": hashlib.sha256(
                        (
                            markdown
                            + json.dumps(table_facts, ensure_ascii=False, sort_keys=True)
                        ).encode("utf-8")
                    ).hexdigest(),
                }
                output_path.write_text(
                    json.dumps(record, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                report.imported_spreadsheets += 1
                report.extracted_characters += len(markdown)
            except Exception as error:
                report.failed_spreadsheets += 1
                report.failures.append({
                    "path": relative_path.as_posix(),
                    "error": f"{type(error).__name__}: {error}",
                })
                logger.error(
                    "Tableur non importé : %s — %s: %s",
                    relative_path, type(error).__name__, error,
                )
            finally:
                if show_progress:
                    spreadsheet_iterator.set_postfix(
                        importés=report.imported_spreadsheets,
                        réutilisés=report.reused_spreadsheets,
                        doublons=report.duplicate_spreadsheets,
                        erreurs=report.failed_spreadsheets,
                        refresh=True,
                    )

        if show_progress:
            spreadsheet_iterator.set_description(
                "Import des tableurs ECO RCI terminé", refresh=True,
            )

        for pattern in ("local_pdf_*.json", "local_sheet_*.json"):
            for stale_path in self.raw_dir.glob(pattern):
                if stale_path.resolve() not in expected_records:
                    stale_path.unlink()
                    report.removed_stale_records += 1

        self._write_report(report)
        logger.info("Import des documents locaux terminé : %s", asdict(report))
        return report


def main() -> None:
    logging.basicConfig(level=logging.INFO)
    report = LocalPDFImporter(load_config()).import_all()
    print(json.dumps(asdict(report), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
