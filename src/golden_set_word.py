"""Construction du golden set de 100 questions a partir du document Word corrige.

Le document source contient les questions dans le corps du fichier et les
metadonnees de reference dans un tableau recapitulatif. Ce module fusionne les
deux representations et resout les chemins documentaires vers les URL reelles
des index economique et Facebook.
"""

from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any
from urllib.parse import quote


_QUESTION_RE = re.compile(r"^(\d+)\.\s+(.+?)\s*$")
_PAGE_RE = re.compile(r"\bpage\s+(\d+)\b", re.IGNORECASE)
_FACEBOOK_RE = re.compile(
    r"(?P<page>[A-Za-z0-9_]+_page)(?:\.txt|\.jsonl)\s*[→>]\s*Publication\s+(?P<number>\d+)",
    re.IGNORECASE,
)
_MULTI_SOURCE_SEPARATOR = " || "

# Les corrections factuelles sont appliquees directement au Word, qui reste
# ainsi l'unique source de verite du golden set.
_REFERENCE_CORRECTIONS: dict[int, dict[str, str]] = {}

_CALCULATED_QUESTIONS = {28, 73, 74}
_MULTI_SOURCE_CALCULATED_QUESTIONS = {77}

_MATCH_STOPWORDS = {
    "a", "au", "aux", "avec", "ce", "ces", "cette", "dans", "de", "des",
    "du", "elle", "en", "est", "et", "ils", "la", "le", "les", "leur",
    "leurs", "lui", "ou", "par", "pour", "qu", "que", "quel", "quelle",
    "quelles", "quels", "qui", "sa", "selon", "son", "sur", "un", "une",
    "etaient", "etait", "presentee", "publication",
}


@dataclass(frozen=True)
class GoldenSetBuildReport:
    source_document: str
    output_path: str
    total_questions: int
    answerable_questions: int
    unanswerable_questions: int
    pdf_questions: int
    facebook_questions: int
    categories: dict[str, int]


def _clean(text: str) -> str:
    return " ".join(str(text).replace("\xa0", " ").split())


def _slugify(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text.casefold())
    ascii_text = "".join(
        character for character in normalized
        if not unicodedata.combining(character)
    )
    return re.sub(r"[^a-z0-9]+", "_", ascii_text).strip("_") or "non_classe"


def _matching_tokens(text: str) -> set[str]:
    normalized = unicodedata.normalize("NFKD", text.casefold())
    normalized = "".join(
        character for character in normalized
        if not unicodedata.combining(character)
    )
    return {
        token for token in re.findall(r"[a-z0-9]+", normalized)
        if token not in _MATCH_STOPWORDS and (len(token) > 1 or token.isdigit())
    }


def _resolve_indexed_facebook_source(
    chunks_path: Path,
    *,
    page_key: str,
    question: str,
    reference_answer: str,
) -> dict[str, Any] | None:
    if not chunks_path.exists():
        return None
    target_tokens = _matching_tokens(f"{question} {reference_answer}")
    if not target_tokens:
        return None
    candidates: dict[str, dict[str, Any]] = {}
    expected_category = f"facebook/{page_key}"
    with chunks_path.open(encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            chunk = json.loads(line)
            if chunk.get("category") != expected_category:
                continue
            parent_id = _clean(chunk.get("parent_id", ""))
            candidate_key = parent_id or _clean(chunk.get("url", ""))
            if candidate_key not in candidates:
                candidates[candidate_key] = chunk

    ranked: list[tuple[float, dict[str, Any]]] = []
    for chunk in candidates.values():
        candidate_text = " ".join(
            str(chunk.get(field, ""))
            for field in ("title", "section", "parent_text", "text")
        )
        candidate_tokens = _matching_tokens(candidate_text)
        overlap = target_tokens & candidate_tokens
        coverage = len(overlap) / len(target_tokens)
        numeric_target = {token for token in target_tokens if token.isdigit()}
        numeric_coverage = (
            len(numeric_target & candidate_tokens) / len(numeric_target)
            if numeric_target else 1.0
        )
        score = 0.85 * coverage + 0.15 * numeric_coverage
        ranked.append((score, chunk))
    if not ranked:
        return None
    score, best = max(ranked, key=lambda item: item[0])
    if score < 0.35:
        return None
    parent_id = _clean(best.get("parent_id", ""))
    post_id_match = re.match(r"facebook_([^_]+)_", parent_id)
    return {
        "url": _clean(best.get("url", "")),
        "title": _clean(best.get("title", "")),
        "post_id": post_id_match.group(1) if post_id_match else "",
        "score": round(score, 6),
    }


def _load_facebook_publication(
    raw_dir: Path,
    page_key: str,
    publication_number: int,
) -> dict[str, Any]:
    raw_path = raw_dir / f"{page_key}.jsonl"
    if not raw_path.exists():
        raise FileNotFoundError(f"Fichier Facebook introuvable : {raw_path}")
    rows = [
        json.loads(line)
        for line in raw_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    if not 1 <= publication_number <= len(rows):
        raise IndexError(
            f"Publication {publication_number} absente de {raw_path.name} "
            f"({len(rows)} publication(s))"
        )
    return rows[publication_number - 1]


def _resolve_source(
    source_locator: str,
    *,
    should_answer: bool,
    facebook_raw_dir: Path,
    facebook_chunks_path: Path,
    question: str,
    reference_answer: str,
) -> dict[str, Any]:
    if _MULTI_SOURCE_SEPARATOR in source_locator:
        resolved_sources = [
            _resolve_source(
                locator.strip(),
                should_answer=should_answer,
                facebook_raw_dir=facebook_raw_dir,
                facebook_chunks_path=facebook_chunks_path,
                question=question,
                reference_answer=reference_answer,
            )
            for locator in source_locator.split(_MULTI_SOURCE_SEPARATOR)
            if locator.strip()
        ]
        urls = list(dict.fromkeys(
            url
            for resolved in resolved_sources
            for url in resolved.get("relevant_urls", [])
        ))
        pages = [
            resolved.get("reference_page")
            for resolved in resolved_sources
            if resolved.get("reference_page") is not None
        ]
        documents = list(dict.fromkeys(
            resolved.get("source_document", "")
            for resolved in resolved_sources
            if resolved.get("source_document")
        ))
        return {
            "relevant_urls": urls,
            "source_kind": "pdf",
            "source_document": " + ".join(documents),
            "reference_page": None,
            "reference_pages": pages,
            "expected_route": "economic",
        }

    page_match = _PAGE_RE.search(source_locator)
    reference_page = int(page_match.group(1)) if page_match else None

    if not should_answer:
        return {
            "relevant_urls": [],
            "source_kind": "none",
            "source_document": "",
            "reference_page": None,
            "expected_route": "economic",
        }

    facebook_match = _FACEBOOK_RE.search(source_locator)
    if facebook_match:
        page_key = facebook_match.group("page")
        publication_number = int(facebook_match.group("number"))
        indexed = _resolve_indexed_facebook_source(
            facebook_chunks_path,
            page_key=page_key,
            question=question,
            reference_answer=reference_answer,
        )
        publication = None
        if indexed is None:
            publication = _load_facebook_publication(
                facebook_raw_dir,
                page_key,
                publication_number,
            )
        post_url = (
            indexed["url"] if indexed is not None
            else _clean(publication.get("post_url", ""))
        )
        if not post_url:
            raise ValueError(
                f"URL absente pour {page_key}, publication {publication_number}"
            )
        return {
            "relevant_urls": [post_url],
            "source_kind": "facebook",
            "source_document": f"{page_key}.jsonl",
            "reference_page": None,
            "expected_route": "facebook",
            "facebook_page_key": page_key,
            "facebook_publication_number": publication_number,
            "facebook_post_id": (
                indexed["post_id"] if indexed is not None
                else _clean(publication.get("post_id", ""))
            ),
            "source_title": indexed["title"] if indexed is not None else "",
            "source_resolution_score": indexed["score"] if indexed is not None else None,
        }

    parts = [_clean(part) for part in re.split(r"\s*[→>]\s*", source_locator)]
    pdf_index = next(
        (index for index, part in enumerate(parts) if part.casefold().endswith(".pdf")),
        None,
    )
    if pdf_index is None:
        raise ValueError(f"Source non reconnue : {source_locator}")
    relative_parts = parts[1 : pdf_index + 1]
    if not relative_parts:
        raise ValueError(f"Chemin PDF incomplet : {source_locator}")
    source_document = relative_parts[-1]
    return {
        "relevant_urls": [
            "/api/source-files/"
            + "/".join(quote(part, safe="") for part in relative_parts)
        ],
        "source_kind": "pdf",
        "source_document": source_document,
        "reference_page": reference_page,
        "expected_route": "economic",
    }


def extract_banque_100_examples(
    document_path: Path,
    *,
    facebook_raw_dir: Path,
    facebook_chunks_path: Path,
) -> list[dict[str, Any]]:
    """Extrait et valide les 100 exemples du fichier Word."""
    try:
        from docx import Document
    except ImportError as error:  # pragma: no cover - message destine au notebook
        raise ImportError(
            "python-docx est requis. Installez les dependances du projet avant "
            "de reconstruire ce golden set."
        ) from error

    document = Document(document_path)
    questions: dict[int, str] = {}
    for paragraph in document.paragraphs:
        match = _QUESTION_RE.match(_clean(paragraph.text))
        if match:
            number = int(match.group(1))
            if number in questions:
                raise ValueError(f"Question {number} presente plusieurs fois")
            questions[number] = match.group(2)

    metadata_table = None
    expected_header = ["n°", "domaine", "statut", "source / chemin", "réponse"]
    for table in document.tables:
        if not table.rows:
            continue
        header = [_clean(cell.text).casefold() for cell in table.rows[0].cells]
        if header == expected_header:
            metadata_table = table
            break
    if metadata_table is None:
        raise ValueError("Tableau recapitulatif des 100 questions introuvable")

    metadata_rows: dict[int, list[str]] = {}
    for row in metadata_table.rows[1:]:
        values = [_clean(cell.text) for cell in row.cells]
        if len(values) != 5 or not values[0].isdigit():
            raise ValueError(f"Ligne de metadonnees invalide : {values}")
        number = int(values[0])
        metadata_rows[number] = values

    expected_numbers = set(range(1, 101))
    if set(questions) != expected_numbers:
        raise ValueError(
            f"Questions attendues 1..100, trouvees : {sorted(questions)}"
        )
    if set(metadata_rows) != expected_numbers:
        raise ValueError(
            f"Metadonnees attendues 1..100, trouvees : {sorted(metadata_rows)}"
        )

    examples: list[dict[str, Any]] = []
    for number in range(1, 101):
        _, domain, status, source_locator, reference_answer = metadata_rows[number]
        reference_correction = _REFERENCE_CORRECTIONS.get(number)
        original_reference_answer = reference_answer
        if reference_correction:
            reference_answer = reference_correction["reference_answer"]
        normalized_status = status.upper().replace("_", "-")
        if normalized_status not in {"IN-BASE", "OUT-OF-BASE"}:
            raise ValueError(f"Statut inconnu pour la question {number} : {status}")
        should_answer = normalized_status == "IN-BASE"
        source = _resolve_source(
            source_locator,
            should_answer=should_answer,
            facebook_raw_dir=facebook_raw_dir,
            facebook_chunks_path=facebook_chunks_path,
            question=questions[number],
            reference_answer=reference_answer,
        )
        examples.append({
            "id": f"banque100-{number:03d}",
            "question_number": number,
            "question": questions[number],
            "reference_answer": reference_answer,
            "original_reference_answer": (
                original_reference_answer if reference_correction else ""
            ),
            "reference_correction_reason": (
                reference_correction["reason"] if reference_correction else ""
            ),
            "relevant_urls": source.pop("relevant_urls"),
            "category": _slugify(domain),
            "domain_label": domain,
            "difficulty": (
                "hard"
                if number in (_CALCULATED_QUESTIONS | _MULTI_SOURCE_CALCULATED_QUESTIONS)
                else "medium"
            ),
            "answer_mode": (
                "calculated_multisource"
                if number in _MULTI_SOURCE_CALCULATED_QUESTIONS
                else "calculated"
                if number in _CALCULATED_QUESTIONS
                else "direct"
                if should_answer
                else "abstention"
            ),
            "split": "banque_100",
            "should_answer": should_answer,
            "answer_keywords": [],
            "source_locator": source_locator,
            "source_status": normalized_status,
            "provenance_document": document_path.name,
            **source,
        })

    answerable = sum(example["should_answer"] for example in examples)
    if answerable != 75:
        raise ValueError(f"75 questions IN-BASE attendues, {answerable} trouvees")
    return examples


def build_banque_100_golden_set(
    document_path: Path,
    output_path: Path,
    *,
    facebook_raw_dir: Path,
    facebook_chunks_path: Path,
) -> GoldenSetBuildReport:
    """Reconstruit le JSONL de facon deterministe depuis le DOCX corrige."""
    document_path = Path(document_path).resolve()
    output_path = Path(output_path).resolve()
    facebook_raw_dir = Path(facebook_raw_dir).resolve()
    facebook_chunks_path = Path(facebook_chunks_path).resolve()
    examples = extract_banque_100_examples(
        document_path,
        facebook_raw_dir=facebook_raw_dir,
        facebook_chunks_path=facebook_chunks_path,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    payload = "\n".join(
        json.dumps(example, ensure_ascii=False) for example in examples
    ) + "\n"
    output_path.write_text(payload, encoding="utf-8")

    categories = Counter(example["domain_label"] for example in examples)
    report = GoldenSetBuildReport(
        source_document=str(document_path),
        output_path=str(output_path),
        total_questions=len(examples),
        answerable_questions=sum(example["should_answer"] for example in examples),
        unanswerable_questions=sum(not example["should_answer"] for example in examples),
        pdf_questions=sum(example["source_kind"] == "pdf" for example in examples),
        facebook_questions=sum(
            example["source_kind"] == "facebook" for example in examples
        ),
        categories=dict(sorted(categories.items())),
    )
    summary_path = output_path.with_name(f"{output_path.stem}_summary.json")
    summary_path.write_text(
        json.dumps(asdict(report), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return report
